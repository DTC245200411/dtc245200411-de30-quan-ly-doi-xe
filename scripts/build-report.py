from pathlib import Path
from docx import Document
from docx.shared import Inches, Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
import json,re,os

ROOT=Path(os.environ.get('FLEET_ROOT',Path(__file__).resolve().parents[1]))
OUT=ROOT/'docs';OUT.mkdir(exist_ok=True)
EVIDENCE=OUT/'evidence';DIAGRAMS=OUT/'diagrams'
doc=Document();section=doc.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=Cm(1.9);section.bottom_margin=Cm(1.8);section.left_margin=Cm(2.4);section.right_margin=Cm(2.0)
section.header_distance=Cm(.7);section.footer_distance=Cm(.7)
section.different_first_page_header_footer=True
styles=doc.styles
for style in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption']:
 st=styles[style];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
 st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Times New Roman')
for st in styles:
 for border in list(st.element.xpath('.//w:pBdr')):
  border.getparent().remove(border)
 for fonts in st.element.xpath('.//w:rFonts'):
  for key in list(fonts.attrib):
   if 'theme' in key.lower():del fonts.attrib[key]
styles['Normal'].font.size=Pt(12)
styles['Normal'].paragraph_format.line_spacing=1.14
styles['Normal'].paragraph_format.space_after=Pt(7)
styles['Normal'].paragraph_format.widow_control=True
styles['Title'].font.size=Pt(22);styles['Title'].font.bold=True
styles['Heading 1'].font.size=Pt(16);styles['Heading 1'].font.bold=True
styles['Heading 1'].paragraph_format.space_before=Pt(0);styles['Heading 1'].paragraph_format.space_after=Pt(12)
styles['Heading 2'].font.size=Pt(15);styles['Heading 2'].font.bold=True
styles['Heading 2'].paragraph_format.space_before=Pt(0);styles['Heading 2'].paragraph_format.space_after=Pt(10)
styles['Caption'].font.size=Pt(10.5);styles['Caption'].font.italic=True
styles['Caption'].paragraph_format.space_after=Pt(9)
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=footer.add_run('Lam Fleet  |  Nguyễn Thị Lam  |  ');run.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.author='Nguyễn Thị Lam';doc.core_properties.last_modified_by='Nguyễn Thị Lam'
doc.core_properties.title='Triển khai và quản trị hệ thống quản lý đội xe Lam Fleet'
doc.core_properties.subject='Đề 30 môn Triển khai và Quản trị Hệ thống Phần mềm'
doc.core_properties.keywords='DTC245200411, CNTT K23D, Vũ Việt Dũng, Fleet, Docker Compose'

def p(text,bold=False):
 para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 para.add_run(text).bold=bold
 return para
def page(title,chapter=None):
 doc.add_page_break()
 if chapter:doc.add_heading(chapter,level=1)
 doc.add_heading(title,level=2)
def code(text):
 para=doc.add_paragraph();para.paragraph_format.space_after=Pt(7)
 for i,line in enumerate(text.splitlines()):
  if i:para.add_run().add_break()
  r=para.add_run(line);r.font.name='Consolas';r.font.size=Pt(9.5)
 return para
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths:
  for c,w in zip(t.columns,widths):c.width=Cm(w)
 for i,h in enumerate(headers):t.rows[0].cells[i].text=h
 for row in rows:
  cells=t.add_row().cells
  for i,value in enumerate(row):cells[i].text=str(value)
 for i,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
  if i==0:repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
  for j,cell in enumerate(row.cells):
   if widths:cell.width=Cm(widths[j])
   cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   props=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'5');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
   props.append(borders)
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8ECE8' if i==0 else ('F7F9F6' if i%2==0 else 'FFFFFF'));props.append(shade)
   margins=OxmlElement('w:tcMar')
   for side,val in [('top','80'),('bottom','80'),('left','100'),('right','100')]:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');margins.append(e)
   props.append(margins)
   for para in cell.paragraphs:
    para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2);para.paragraph_format.line_spacing=1.05
    for run in para.runs:run.font.size=Pt(10.5);run.bold=i==0
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
 return t
def image(path,caption,width=6.3,maxheight=5):
 assert path.is_file(),str(path)
 with Image.open(path) as im:w,h=im.size
 width=min(width,maxheight*w/h)
 para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.space_after=Pt(3)
 shape=para.add_run().add_picture(str(path),width=Inches(width))
 shape._inline.docPr.set('descr',caption)
 cap=doc.add_paragraph(caption,style='Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
def cover(text,size=14,bold=False,after=10):
 para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.space_after=Pt(after)
 r=para.add_run(text);r.font.size=Pt(size);r.bold=bold
 return para

# 1 Bìa
cover('TRƯỜNG ĐẠI HỌC CÔNG NGHỆ THÔNG TIN\nVÀ TRUYỀN THÔNG',14,True,6)
cover('KHOA CÔNG NGHỆ THÔNG TIN',14,True,37)
para=doc.add_paragraph('BÁO CÁO BÀI TẬP LỚN',style='Title');para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.space_after=Pt(18)
cover('TRIỂN KHAI VÀ QUẢN TRỊ\nHỆ THỐNG PHẦN MỀM',18,True,30)
cover('ĐỀ 30',16,True,10)
cover('HỆ THỐNG QUẢN LÝ ĐỘI XE',18,True,12)
cover('LAM FLEET',22,True,39)
for text in ['Sinh viên thực hiện  Nguyễn Thị Lam','Mã số sinh viên  DTC245200411','Lớp  CNTT K23D','Giảng viên hướng dẫn  Vũ Việt Dũng']:cover(text,14,False,9)
cover('Thái Nguyên tháng 10 năm 2026',13,False,0).paragraph_format.space_before=Pt(50)

# 2 Mục lục
page('MỤC LỤC')
toc=[('Chương 1 Tổng quan kiến trúc và môi trường',3),('1.1 Mục tiêu và phạm vi',3),('1.2 Kiến trúc triển khai WSL2',4),('1.3 Quản lý mã nguồn trên GitHub',5),('Chương 2 Ứng dụng Database và Reverse Proxy',6),('2.1 Mô hình dữ liệu và ràng buộc',6),('2.2 PostgreSQL và pgAdmin',7),('2.3 Hồ sơ xe và tài xế',8),('2.4 Điều phối và chống trùng lịch',9),('2.5 Vòng đời bảo dưỡng',10),('2.6 Nginx và cách khởi động',11),('Chương 3 Giám sát và log tập trung',12),('3.1 Prometheus và exporter',12),('3.2 Dashboard Grafana',13),('3.3 Loki Promtail và LogQL',14),('Chương 4 Hardening và đánh giá',15),('4.1 Biện pháp gia cố',15),('4.2 Kết quả kiểm thử',16),('4.3 Sao lưu và phục hồi',17),('Kết luận kịch bản demo và tài liệu tham khảo',18)]
for title,n in toc:
 para=doc.add_paragraph();para.paragraph_format.space_after=Pt(4);para.paragraph_format.tab_stops.add_tab_stop(Cm(16),WD_ALIGN_PARAGRAPH.RIGHT)
 r=para.add_run(title+'\t'+str(n));r.bold=title.startswith('Chương');r.font.size=Pt(12)
p('Báo cáo trình bày hệ thống đã triển khai bằng Docker Engine trong Ubuntu WSL2. Các ảnh, số liệu và kết quả kiểm tra là minh chứng của project de30lam; hồ sơ xe và tài xế trong dữ liệu khởi tạo là nội dung minh họa.')

# 3
page('1.1 Mục tiêu và phạm vi','CHƯƠNG 1 TỔNG QUAN KIẾN TRÚC VÀ MÔI TRƯỜNG')
p('Lam Fleet phục vụ một quản trị viên quản lý phương tiện, tài xế, chuyến đi và bảo dưỡng. Người quản trị xem tình trạng đội xe, xếp lịch theo khoảng thời gian và lưu kết quả công việc. Mục tiêu vận hành là triển khai đồng bộ ứng dụng, database, công cụ quản trị, reverse proxy, giám sát và log tập trung bằng Docker Compose.')
image(DIAGRAMS/'usecase.png','Hình 1 Các nhóm chức năng của quản trị viên',width=4.6,maxheight=4.8)
p('Phạm vi bài thực hành gồm một đơn vị quản lý đội xe, một tài khoản quản trị chính và dữ liệu mẫu. Hệ thống không tích hợp GPS, thanh toán, định tuyến bản đồ hoặc ứng dụng di động. Giới hạn này giúp tập trung kiểm chứng khả năng điều phối và quản trị hạ tầng.')
p('Tiêu chí hoàn thành là website thao tác được qua Nginx ở cổng 8080, pgAdmin kết nối database fleet, Grafana hiển thị metrics thực, truy vấn LogQL có dữ liệu và các biện pháp hardening có kết quả kiểm tra. Báo cáo đi kèm mã nguồn và hướng dẫn tự chạy lại.')

# 4
page('1.2 Kiến trúc triển khai WSL2')
p('Chrome chạy trên Windows, còn toàn bộ dịch vụ nằm trong Docker Engine của Ubuntu WSL2. Compose project de30lam có network và volume riêng; việc triển khai không thay đổi các bài thực hành khác đang chạy. Nguồn chính đặt trong filesystem Linux để giữ quyền truy cập và hiệu năng mount.')
image(DIAGRAMS/'architecture.png','Hình 2 Kiến trúc triển khai ứng dụng và database',width=4.7,maxheight=5.05)
table(['Dịch vụ','Địa chỉ trên Windows','Cách truy cập'],[['Website','localhost:8080','Qua Nginx'],['pgAdmin','localhost:8091','admin@lam.edu.vn'],['Grafana','localhost:3005','admin'],['Prometheus','localhost:9094','Trang Targets và truy vấn']],widths=[3.3,6.0,7.0])
p('PostgreSQL và app không publish cổng ra host. Các cổng giao diện chỉ bind 127.0.0.1. Khi demo từ Windows, cần giữ phiên Ubuntu hoạt động để dịch vụ và chuyển tiếp localhost không bị gián đoạn.')

# 5
page('1.3 Quản lý mã nguồn trên GitHub')
p('Repository riêng tư thuộc tài khoản GitHub DTC245200411, trùng mã số sinh viên. Trước khi tạo repository, GitHub MCP đã trả đúng tài khoản này; Git Credential Manager trên Windows và kết nối Git trong WSL cũng được kiểm chứng qua GitHub API. Git dùng tác giả Nguyễn Thị Lam và email noreply của đúng tài khoản.')
p('Repository bàn giao có địa chỉ https://github.com/DTC245200411/dtc245200411-de30-quan-ly-doi-xe. Nhánh bàn giao trên GitHub là main. Mỗi mốc có commit riêng phản ánh đúng phần việc được đưa vào mã nguồn, thay vì tạo lịch sử rỗng để đủ số lượng.')
table(['Mốc','Nội dung thực hiện','Minh chứng'],[['1','Ứng dụng Fleet, PostgreSQL, pgAdmin, Nginx, xác thực và kiểm thử nghiệp vụ','Commit 70b2173'],['2','Prometheus, Grafana, exporter và số liệu container thuộc de30lam','Commit 8d2fa4a'],['3','Loki, Promtail, LogQL, hardening, backup và tài liệu bàn giao','Commit hoàn thiện trên main']],widths=[1.3,10.5,4.5])
p('Thư mục app chứa schema, migration, domain, giao diện và kiểm thử. infra chứa cấu hình PostgreSQL, pgAdmin, Nginx và xoay vòng log. observability chứa các datasource, dashboard và cấu hình thu thập. scripts cung cấp thao tác vận hành; docs giữ sơ đồ và minh chứng.')
p('Các tệp .env, .local, secret và bản backup không được commit. .env.example chỉ khai báo cổng mẫu, không có mật khẩu. Script chuẩn bị sinh mật khẩu riêng khi chạy lần đầu và giữ lại khi chạy tiếp. Dependency được khóa bằng package-lock.json; image Docker dùng tag phiên bản cụ thể.')
code('git log --oneline\ngit status --short\ngit remote -v')
p('Người clone repository riêng tư phải có quyền truy cập. README là tài liệu vận hành đầu tiên; sơ đồ có nguồn DOT để chỉnh sửa và tái tạo.')

# 6
page('2.1 Mô hình dữ liệu và ràng buộc','CHƯƠNG 2 ỨNG DỤNG DATABASE VÀ REVERSE PROXY')
p('PostgreSQL lưu bốn nhóm nghiệp vụ chính. Mỗi chuyến bắt buộc tham chiếu một xe và một tài xế; mỗi phiếu bảo dưỡng tham chiếu một xe. Một xe hoặc tài xế có thể chưa có lịch sử, vì vậy phía tập lịch có lực lượng 0..n.')
image(DIAGRAMS/'erd.png','Hình 3 ERD vật lý rút gọn của bốn bảng nghiệp vụ',width=6.5,maxheight=4.9)
table(['Bảng bổ trợ','Vai trò','Ràng buộc chính'],[['admins','Tài khoản quản trị','username duy nhất, mật khẩu băm scrypt'],['sessions','Phiên đăng nhập','sid là PK; sess là JSON; không có FK đến admins'],['seed_versions','Đánh dấu seed đã áp dụng','version là PK; seed không chạy lặp']],widths=[3.0,5.0,8.3])
p('Các mã xe, tài xế, chuyến và phiếu đều duy nhất; biển số và số giấy phép cũng duy nhất. start_at và end_at dùng timestamptz, với end_at lớn hơn start_at. Số kilomet và chi phí không âm. Chỉ mục hỗ trợ tra lịch mở theo xe, tài xế và thời gian; version hỗ trợ phát hiện cập nhật từ biểu mẫu cũ.')

# 7
page('2.2 PostgreSQL và pgAdmin')
p('Database fleet chạy trên PostgreSQL 17. Role fleet_owner phục vụ khởi tạo và sao lưu; fleet_migrate tạo schema và seed; fleet_app thực hiện nghiệp vụ theo quyền được cấp; fleet_exporter chỉ phục vụ giám sát. Ứng dụng không có quyền tạo bảng hoặc xóa lịch sử chuyến và bảo dưỡng.')
image(EVIDENCE/'13-pgadmin.png','Hình 4 pgAdmin kết nối thành công tới database fleet',width=6.5,maxheight=4.6)
p('pgAdmin truy cập tại cổng 8091 với email admin@lam.edu.vn. Server Lam Fleet PostgreSQL được khai báo sẵn, trỏ postgres:5432, database fleet, tài khoản fleet_migrate. Mật khẩu đăng nhập pgAdmin và mật khẩu kết nối database là hai secret khác nhau; cả hai chỉ lưu trong .local/secrets.')
p('Migration và seed chạy trong transaction, có advisory lock để không khởi tạo đồng thời. Phiên bản seed chỉ được đánh dấu khi toàn bộ dữ liệu ghi thành công. Bộ dữ liệu gồm 8 xe, 6 tài xế, 10 chuyến và 5 phiếu bảo dưỡng; có xe ngừng sử dụng và nhiều trạng thái lịch để demo.')
code('SELECT status, count(*) FROM trips GROUP BY status;\nSELECT status, count(*) FROM maintenance GROUP BY status;')

# 8
page('2.3 Hồ sơ xe và tài xế')
p('Trang Đội xe tổng hợp mã xe, tên, biển số, loại, số kilomet và tình trạng vận hành. Trạng thái sẵn sàng, đang chạy hoặc đang bảo dưỡng được suy ra từ hồ sơ và công việc thực tế đang mở, giúp tránh lưu hai nguồn trạng thái mâu thuẫn. Người dùng tìm theo mã, biển số hoặc tên xe.')
image(EVIDENCE/'03-vehicles.png','Hình 5 Danh sách phương tiện và trạng thái vận hành',width=6.5,maxheight=4.7)
p('Hồ sơ tài xế có mã duy nhất, họ tên, điện thoại và số giấy phép. Người quản trị được sửa thông tin và chuyển ngừng hoạt động khi không còn lịch mở. Trang chi tiết liên kết tới các chuyến đã thực hiện; số điện thoại và giấy phép trong seed chỉ là dữ liệu minh họa.')
p('Khi cập nhật xe, hệ thống không chấp nhận kilomet giảm. Nếu xe đang thực hiện công việc, kilomet chỉ được thay đổi khi hoàn tất công việc đó. Hồ sơ được khóa trong transaction, đồng thời kiểm tra version để một biểu mẫu cũ không ghi đè thay đổi mới. Hệ thống giữ lịch sử và không cung cấp thao tác xóa vật lý hồ sơ nghiệp vụ.')

# 9
page('2.4 Điều phối và chống trùng lịch')
p('Một lịch được biểu diễn bằng khoảng nửa mở [bắt đầu, kết thúc). Hai lịch tiếp giáp được chấp nhận; hai lịch có giao nhau trên cùng xe hoặc cùng tài xế bị từ chối. Kiểm tra xe bao gồm cả chuyến đi và bảo dưỡng ở trạng thái dự kiến hoặc đang thực hiện.')
image(DIAGRAMS/'workflow.png','Hình 6 Transaction và thứ tự khóa khi ghi lịch',width=5.8,maxheight=5.3)
p('Các thao tác khóa xe theo ID tăng dần trước, khóa tài xế sau và cuối cùng khóa lịch khi sửa. Sau khi lấy khóa, hệ thống đọc lại dữ liệu và kiểm tra giao nhau trước khi ghi. Hai yêu cầu cùng đặt một tài nguyên vì thế có tối đa một yêu cầu thành công; yêu cầu còn lại nhận HTTP 409 cùng thông báo tiếng Việt.')
p('Chuyến dự kiến có thể sửa hoặc hủy. Khi bắt đầu, hệ thống kiểm tra thêm công việc đang chạy, kể cả trường hợp công việc đó đã quá giờ dự kiến. Khi hoàn tất, kilomet cuối không được giảm và được cập nhật về hồ sơ xe. Bài không tự động điều chỉnh lịch khác khi một chuyến thực tế kéo dài.')

# 10
page('2.5 Vòng đời bảo dưỡng')
p('Phiếu bảo dưỡng lưu mã phiếu, xe, nội dung, đơn vị thực hiện và khoảng thời gian dự kiến. Phiếu chưa bắt đầu có thể chỉnh sửa hoặc hủy. Khi bắt đầu, xe không được đồng thời có chuyến hoặc bảo dưỡng khác đang thực hiện. Trang tổng quan đưa các lịch mở lên để người quản trị xử lý, đánh dấu lịch dự kiến đã quá giờ bắt đầu.')
image(EVIDENCE/'08-maintenance-detail.png','Hình 7 Ghi nhận hoàn tất bảo dưỡng và chi phí thực tế',width=6.5,maxheight=4.8)
p('Hoàn tất bảo dưỡng cần số kilomet hợp lệ, chi phí VND nguyên không âm và kết quả thực hiện. Thay đổi trạng thái và cập nhật kilomet xe nằm trong cùng transaction. Phiếu đã hoàn tất giữ nguyên lịch sử để tra cứu; tổng chi phí trên dashboard chỉ cộng các phiếu đã hoàn tất.')
p('Cảnh báo quá hạn dựa trên giờ bắt đầu dự kiến, không phải cảnh báo định kỳ theo kilomet. Việc lập kế hoạch lần bảo dưỡng kế tiếp do quản trị viên thực hiện. Đây là phạm vi đủ để minh họa nghiệp vụ quản lý và các sự kiện log trong bài thực hành.')

# 11
page('2.6 Nginx và cách khởi động')
p('Nginx là cổng truy cập website duy nhất ở localhost:8080 và chuyển tiếp tới app:3000 trong mạng frontend. Cấu hình có giới hạn kích thước biểu mẫu, timeout và các security header. Các endpoint /metrics và /readyz bị chặn ở cổng website; exporter truy cập qua mạng monitoring riêng.')
image(EVIDENCE/'02-dashboard.png','Hình 8 Trang tổng quan Lam Fleet được phục vụ qua Nginx',width=6.5,maxheight=4.6)
code('sh scripts/start.sh\nsh scripts/test.sh\nsh scripts/stop.sh')
p('Script start chuẩn bị secret, kiểm tra Compose, build image, đợi dịch vụ sẵn sàng và kiểm tra cấu hình Nginx. Có thể dùng BASE_ONLY=1 để chỉ chạy ứng dụng và database. PowerShell cung cấp wrapper gọi đúng Ubuntu WSL2. Khi dùng trực tiếp lệnh Linux, giữ cửa sổ Ubuntu mở trong suốt buổi demo.')
p('HTTP localhost cùng security headers đáp ứng lựa chọn triển khai của đề. Khi đưa ra Internet cần bổ sung chứng chỉ HTTPS hợp lệ, cookie Secure và kiểm soát các cổng quản trị; cấu hình demo hiện tại không được mô tả là đã triển khai HTTPS.')

# 12
page('3.1 Prometheus và exporter','CHƯƠNG 3 GIÁM SÁT VÀ LOG TẬP TRUNG')
p('Prometheus thu thập số liệu mỗi 15 giây từ năm job: prometheus, lam-fleet, nginx, postgres và cadvisor. App công bố metrics HTTP và trạng thái nghiệp vụ. Exporter Nginx đọc stub_status nội bộ; exporter PostgreSQL dùng role giám sát. cAdvisor cung cấp số liệu container trên Docker host.')
image(EVIDENCE/'10-prometheus.png','Hình 9 Các target của project Lam Fleet trong Prometheus',width=6.5,maxheight=4.9)
p('Kết quả kiểm tra xác nhận đủ 5 target UP, nginx_up và pg_up bằng 1. CPU và RAM phải có nhãn Compose project de30lam và đúng service app, nginx, postgres; số liệu chỉ thuộc Docker host hoặc bài khác không được coi là đạt. Gauge tổng xe đọc trực tiếp PostgreSQL và bằng 8.')
p('Biểu đồ rate cần ít nhất hai lần thu thập. Các bộ đếm HTTP và thao tác bắt đầu lại khi app restart; số chuyến và phiếu trong database vẫn giữ. Dữ liệu Prometheus được giữ 7 ngày và giới hạn dung lượng theo cấu hình demo.')

# 13
page('3.2 Dashboard Grafana')
p('Grafana được provision tự động với datasource Prometheus và Loki. Dashboard Lam Fleet Đội xe và vận hành nằm trong thư mục Đề 30 Nguyễn Thị Lam, gồm 14 panel metrics và 3 panel log. Người dùng không phải tạo thủ công dashboard sau khi clone dự án.')
image(EVIDENCE/'11-grafana-metrics.png','Hình 10 Dashboard với số liệu CPU RAM HTTP và database',width=6.5,maxheight=4.8)
p('Các panel đầu hiển thị số target đang hoạt động, tổng xe, chuyến đang thực hiện và bảo dưỡng dự kiến. Nhóm biểu đồ theo dõi CPU/RAM, lưu lượng HTTP theo mã phản hồi, p95 độ trễ, kết nối Nginx, kết nối và giao dịch PostgreSQL. Nhóm nghiệp vụ thể hiện trạng thái chuyến, bảo dưỡng và số yêu cầu xung đột bị từ chối.')
p('Truy vấn container luôn lọc project de30lam. Dashboard phục vụ việc đối chiếu tác động của thao tác demo lên hệ thống, không thay thế sổ lịch trình trong database. Alert target mất kết nối được khai báo trong Prometheus; bài không cấu hình gửi email hay thông báo ra dịch vụ bên ngoài.')

# 14
page('3.3 Loki Promtail và LogQL')
p('Ứng dụng, Nginx và PostgreSQL ghi log ra volume riêng. Promtail mount log ở chế độ chỉ đọc, lưu vị trí đọc trên volume độc lập rồi đẩy vào Loki. Grafana dùng LogQL để tìm sự kiện. Loki giữ log 7 ngày; logrotate quản lý file nhằm hạn chế tăng dung lượng trên máy demo.')
image(EVIDENCE/'12-grafana-logs.png','Hình 11 Truy vấn log nghiệp vụ và lỗi database trên Grafana',width=6.5,maxheight=3.55)
for title,query in [('Lỗi HTTP','{job="nginx", project="de30lam"}\n  | json | status >= 400'),('Thao tác nghiệp vụ','{job="app", project="de30lam"} | json\n  | event=~"(trip|maintenance)_(scheduled|started|completed|cancelled)"'),('Lỗi PostgreSQL','{job="postgres", project="de30lam"} | json\n  | error_severity=~"ERROR|FATAL|PANIC"')]:
 para=doc.add_paragraph(title);para.paragraph_format.space_after=Pt(2);para.runs[0].bold=True;code(query)
logging=json.loads((EVIDENCE/'logging.json').read_text())
p('Lần kiểm chứng thu được lần lượt '+', '.join(str(r['events']) for r in logging['results'])+' sự kiện theo giới hạn truy vấn. Mật khẩu, cookie và nội dung form không được đưa vào log. Promtail được giữ theo yêu cầu bài; Grafana xác nhận công cụ đã hết hỗ trợ ngày 02/03/2026, nên hướng nâng cấp là Alloy [4].')

# 15
page('4.1 Biện pháp gia cố','CHƯƠNG 4 HARDENING VÀ ĐÁNH GIÁ')
p('Gia cố được áp dụng ở cả ứng dụng, container và database. Script check-hardening.py đọc cấu hình container đang chạy để xác nhận quyền người dùng, filesystem, capability, network và cách publish cổng, thay vì chỉ dựa trên khai báo trong báo cáo.')
table(['Lớp','Biện pháp','Kết quả đối chiếu'],[['Ứng dụng','Mật khẩu băm scrypt, CSRF, giới hạn đăng nhập, kiểm tra dữ liệu','HTTP test đạt'],['Phiên','Cookie HttpOnly và SameSite; đổi mật khẩu làm hết hiệu lực phiên cũ','HTTP test đạt'],['App và Nginx','Non-root, read-only rootfs, cap_drop ALL, no-new-privileges','Container inspect đạt'],['PostgreSQL','Tách role; app không có CREATE schema hoặc DELETE lịch sử','Kiểm tra quyền đạt'],['Network','Database nội bộ; không publish app, DB, Loki hoặc exporter','Container inspect đạt'],['Cổng giao diện','Website và công cụ chỉ bind 127.0.0.1','Port binding đạt'],['Secret','Sinh riêng, thư mục 0700, không có trong Git','Kiểm tra filesystem và Git đạt'],['Log','Mount log Promtail chỉ đọc; không mount Docker socket','Container inspect đạt']],widths=[2.5,8.6,5.2])
p('Các ngoại lệ được ghi rõ: cAdvisor cần privileged và truy cập Docker host; storage-init dùng root tạm thời để đặt quyền volume; PostgreSQL bootstrap bắt đầu bằng root rồi chuyển sang postgres; logrotate có một số capability phục vụ file log. Vì vậy không khẳng định tất cả container đều non-root hoặc không có đặc quyền.')
p('Các biện pháp trên phù hợp môi trường demo nội bộ. Chưa thực hiện kiểm thử xâm nhập toàn diện hay đánh giá tải lớn. Trước khi cung cấp dịch vụ công khai cần bổ sung HTTPS, chính sách truy cập quản trị, cập nhật công cụ log còn hỗ trợ và kế hoạch sao lưu ngoài máy.')

# 16
page('4.2 Kết quả kiểm thử')
p('Bộ kiểm thử dùng Node test runner, gửi HTTP thật qua Nginx và đối chiếu PostgreSQL. Mỗi lần chạy tạo tài khoản QA và dữ liệu có tiền tố riêng, sau đó dọn dữ liệu thử. Mật khẩu của tài khoản admin chính không bị thay đổi. Kết quả lần bàn giao là 29/29 trường hợp đạt, không có trường hợp bỏ qua.')
table(['Nhóm kiểm thử','Số ca','Kết quả'],[['Xác thực CSRF headers cookie và đổi mật khẩu',5,'Đạt'],['Màn hình và tìm kiếm tham số hóa',2,'Đạt'],['Dữ liệu trạng thái quyền DB và seed',6,'Đạt'],['Tạo sửa hủy và hoàn tất chuyến',5,'Đạt'],['Bảo dưỡng chi phí và công việc đang chạy',3,'Đạt'],['Yêu cầu đồng thời',4,'Đạt'],['Kiểm tra số mã và thời gian',4,'Đạt'],['Tổng cộng',29,'29 đạt 0 lỗi']],widths=[11.0,2.0,3.3])
p('Các ca đồng thời gửi hai yêu cầu cùng đặt xe, cùng đặt tài xế, một chuyến cạnh tranh với một phiếu bảo dưỡng và hai lần bắt đầu cùng một lịch. Mỗi trường hợp có đúng một yêu cầu thành công; yêu cầu còn lại bị từ chối, không ghi dữ liệu một phần. Lịch tiếp giáp vẫn được chấp nhận.')
code('# tests 29\n# pass 29\n# fail 0\n# skipped 0')
p('Kiểm thử trình duyệt Chrome chụp các màn hình và xác nhận không có lỗi JavaScript ở website. Màn hình desktop và viewport 390 px không gây tràn ngang toàn trang; bảng dài có vùng cuộn riêng. npm audit trên dependency đã khóa báo 0 lỗ hổng tại thời điểm triển khai. Những kết quả này kiểm chứng phạm vi đã thử, không phải cam kết không còn mọi lỗi bảo mật.')

# 17
page('4.3 Sao lưu và phục hồi')
p('Script backup dùng pg_dump dạng custom để lưu schema, dữ liệu, ownership và ACL. Mỗi archive đi kèm SHA-256; thư mục backup không nằm trong Git và được tạo với quyền hạn chế. File có dữ liệu tài khoản và phiên nên không được chia sẻ như minh chứng công khai.')
code('sh scripts/backup.sh\nsh scripts/restore.sh --verify-current .local/backups/<tep>.dump')
p('Chế độ verify-current kiểm tra checksum, tạo database tạm, phục hồi archive không ghi đè database đang dùng, rồi so sánh fingerprint của bốn bảng nghiệp vụ. Database tạm được dọn sau kiểm tra. So sánh này xác nhận nội dung dữ liệu nghiệp vụ khớp, không chỉ đếm số dòng.')
table(['Bảng nghiệp vụ','Số dòng phục hồi','Đối chiếu'],[['vehicles',8,'Khớp'],['drivers',6,'Khớp'],['trips',10,'Khớp'],['maintenance',5,'Khớp']],widths=[7.0,4.5,4.8])
restore=(EVIDENCE/'restore.txt').read_text(encoding='utf-8')
match=re.search(r'"fingerprint"\s*:\s*"([a-f0-9]+)"',restore)
if match:p('Fingerprint đã đối chiếu: '+match.group(1)+'. Kết quả ghi nhận PASS restored business rows match current database exactly trong minh chứng restore.txt.')
p('Chế độ --apply phục hồi thay dữ liệu hiện có. Script kiểm tra archive bằng database tạm trước, yêu cầu nhập PHUC HOI LAM FLEET, dừng app/Nginx rồi mới phục hồi. Nếu thất bại, ứng dụng giữ trạng thái dừng để kiểm tra; không tiếp tục phục vụ với dữ liệu chưa rõ trạng thái. Bài bàn giao đã thử phục hồi vào database tạm, không thực hiện ghi đè database demo.')
p('Dừng hệ thống thông thường giữ named volume. Khởi động lại chạy migration và kiểm tra seed_versions, nên không nhân đôi dữ liệu hoặc đặt lại mật khẩu. Không sử dụng lệnh prune toàn máy hoặc xóa volume của bài khác để xử lý lỗi.')

# 18
page('KẾT LUẬN KỊCH BẢN DEMO VÀ TÀI LIỆU THAM KHẢO')
p('Lam Fleet đã tích hợp ứng dụng quản lý đội xe, PostgreSQL, pgAdmin, Nginx, Prometheus, Grafana và Loki/Promtail bằng Docker Compose. Kết quả có 29 kiểm thử đạt, 5 target UP, metrics đúng project, ba truy vấn log có dữ liệu và bản backup phục hồi khớp dữ liệu nghiệp vụ. Mã nguồn và cấu hình được bàn giao trên repository riêng của sinh viên.')
table(['Bước','Thao tác demo','Kết quả cần quan sát'],[['1','Mở localhost:8080 và đăng nhập','Dashboard có 8 xe và trạng thái vận hành'],['2','Tạo xe/tài xế thử và lập chuyến','Lịch lưu thành công và xuất hiện trong danh sách'],['3','Lập lịch chồng xe hoặc tài xế','Thông báo xung đột, dữ liệu cũ giữ nguyên'],['4','Bắt đầu rồi hoàn tất chuyến/bảo dưỡng','Trạng thái và kilomet cập nhật đồng bộ'],['5','Mở pgAdmin Prometheus và Grafana','Database fleet, 5 target UP, metrics và log'],['6','Chạy kiểm thử và verify backup','Kết quả PASS và dữ liệu phục hồi khớp']],widths=[1.2,7.8,7.3])
p('Hướng phát triển gồm phân quyền nhiều người dùng, nhắc bảo dưỡng theo kilomet, xử lý chuyến kéo dài ảnh hưởng lịch sau, HTTPS và chuyển Promtail sang Grafana Alloy. Bản hiện tại không có GPS, thanh toán hoặc kiểm chứng tải lớn. Khi cần chạy lại, bắt đầu từ README và kịch bản demo trong thư mục docs.')
doc.add_heading('Tài liệu tham khảo',level=3)
for text in [
 '[1] Đề tài thực hành môn Triển khai và Quản trị Hệ thống Phần mềm K23, yêu cầu chung và đề 30, tài liệu PDF được cung cấp.',
 '[2] Docker Compose. https://docs.docker.com/compose/',
 '[3] PostgreSQL 17 Explicit Locking. https://www.postgresql.org/docs/17/explicit-locking.html',
 '[4] Grafana Promtail agent. https://grafana.com/docs/loki/latest/send-data/promtail/',
 '[5] Prometheus Overview. https://prometheus.io/docs/introduction/overview/'
]:
 para=doc.add_paragraph(text);para.paragraph_format.space_after=Pt(4)
 for run in para.runs:run.font.size=Pt(10)

target=OUT/'Bao_Cao_De30_Nguyen_Thi_Lam.docx';doc.save(target)
print('Created',target)
print('Explicit pages',1+sum(1 for e in doc.element.iter(qn('w:br')) if e.get(qn('w:type'))=='page'))
